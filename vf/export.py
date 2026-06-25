"""Export des résultats : CSV, Excel, rapports marché et synthèse top deals."""

# --- catalog.py ---

import csv
import json
import time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from .market import CatalogItem

LISTING_EXPORT_COLUMNS: List[str] = [
    "listing_key",
    "id",
    "platform",
    "platform_item_id",
    "title",
    "brand",
    "model",
    "model_key",
    "product_model_label",
    "price_eur",
    "total_eur",
    "currency",
    "condition",
    "description",
    "description_excerpt",
    "photo_url",
    "photo_urls",
    "url",
    "seller_name",
    "seller_rating",
    "seller_feedback",
    "shipping_available",
    "hand_delivery_only",
    "delivery_mode",
    "engagement",
    "market_category",
    "search_intent",
    "listed_at",
    "posted_at",
    "listing_age_days",
    "listing_age_label",
    "age_text",
    "age_hours",
    "scraped_at",
    "final_score",
    "composite_score",
    "score_composite",
    "deal_tier",
    "analysis_detail",
    "dropship_risk_score",
    "comparison_insight",
    "rank_score",
    "display_rank_score",
    "market_discount_pct",
    "discount_vs_new_pct",
    "buyer_fit_score",
    "buyer_fit_grade",
    "quality_tier",
    "spec_tier",
    "pack_score",
    "is_accessory",
    "accessory_kind",
    "risk_flags",
    "triage_decision",
    "triage_reasons",
    "ai_specs_summary",
    "ai_verdict",
    "ai_vision_note",
    "spec_summary",
]


def listing_key(platform: str, platform_item_id: str) -> str:
    plat = (platform or "unknown").strip().lower()
    pid = str(platform_item_id or "").strip()
    return f"{plat}:{pid}" if pid else plat


def _as_float(v: Any) -> Optional[float]:
    try:
        if v is None or v == "":
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def normalize_listing_row(
    row: Dict[str, Any],
    *,
    search_intent: str = "",
    market_category: str = "",
) -> Dict[str, Any]:
    """Normalise les noms de champs pour comparaison multi-plateformes."""
    plat = str(row.get("platform") or "vinted").lower()
    pid = str(row.get("platform_item_id") or row.get("id") or "").strip()
    out = dict(row)
    out["platform"] = plat
    out["platform_item_id"] = pid
    out["listing_key"] = listing_key(plat, pid)
    out["id"] = str(row.get("id") or pid or out["listing_key"])
    out["price_eur"] = _as_float(row.get("price_eur") or row.get("price"))
    out["total_eur"] = _as_float(row.get("total_eur") or row.get("total_price") or row.get("price_eur"))
    out["seller_name"] = str(row.get("seller_name") or row.get("seller_username") or "").strip()
    out["seller_feedback"] = row.get("seller_feedback") or row.get("seller_feedback_count")
    if not out.get("description_excerpt") and out.get("description"):
        out["description_excerpt"] = str(out["description"])[:280]
    photos = row.get("photo_urls")
    if isinstance(photos, list):
        out["photo_urls"] = "|".join(str(u) for u in photos if u)[:2000]
    elif photos:
        out["photo_urls"] = str(photos)[:2000]
    if search_intent:
        out["search_intent"] = search_intent[:200]
    if market_category:
        out["market_category"] = market_category
    if not out.get("scraped_at"):
        out["scraped_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
    try:
        from .scoring import enrich_row_freshness

        enrich_row_freshness(out)
    except Exception:
        pass
    try:
        from .scoring import enrich_row_landed_cost, is_discontinued_product

        enrich_row_landed_cost(out)
        if is_discontinued_product(out):
            out["is_discontinued"] = True
    except Exception:
        pass
    try:
        from .engine import enrich_row_delivery

        enrich_row_delivery(out)
    except Exception:
        pass
    return out


def normalize_catalog_item(
    item: CatalogItem,
    *,
    search_intent: str = "",
    market_category: str = "",
) -> Dict[str, Any]:
    return normalize_listing_row(
        item.to_row(),
        search_intent=search_intent,
        market_category=market_category,
    )


def normalize_catalog_items(
    items: Sequence[CatalogItem],
    *,
    search_intent: str = "",
    market_category: str = "",
) -> List[Dict[str, Any]]:
    return [
        normalize_catalog_item(it, search_intent=search_intent, market_category=market_category)
        for it in items
    ]


def export_listings_csv(
    path: Path,
    rows: Iterable[Dict[str, Any]],
    *,
    columns: Optional[Sequence[str]] = None,
) -> int:
    """Export CSV UTF-8-BOM — colonnes catalogue + champs extra."""
    rows_list = [normalize_listing_row(r) for r in rows]
    if not rows_list:
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    base_cols = list(columns or LISTING_EXPORT_COLUMNS)
    extra: List[str] = []
    for r in rows_list:
        for k in r:
            if k not in base_cols and not str(k).startswith("_"):
                extra.append(k)
    fields = base_cols + [k for k in extra if k not in base_cols]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows_list:
            row_out = dict(r)
            if isinstance(row_out.get("photo_urls"), list):
                row_out["photo_urls"] = "|".join(str(u) for u in row_out["photo_urls"])
            w.writerow(row_out)
    return len(rows_list)


def export_listings_json(path: Path, rows: Sequence[Dict[str, Any]], *, meta: Optional[Dict[str, Any]] = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "meta": meta or {},
        "listings": [normalize_listing_row(r) for r in rows],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


# --- excel_export.py ---

import statistics
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

# French column headers for main offers sheet
OFFRES_COLUMNS: List[tuple[str, str]] = [
    ("id", "ID"),
    ("title", "Titre"),
    ("price", "Prix"),
    ("total_price", "Prix total"),
    ("platform", "Plateforme"),
    ("url", "URL"),
    ("photo_url", "Photo"),
    ("model", "Modèle"),
    ("variant", "Variante"),
    ("storage_gb", "Stockage (Go)"),
    ("controllers_count", "Manettes"),
    ("games_count", "Nb jeux"),
    ("games_list", "Liste jeux"),
    ("premium_games", "Jeux premium"),
    ("condition", "État"),
    ("has_box", "Boîte"),
    ("has_hdmi", "HDMI"),
    ("has_power_cable", "Alimentation"),
    ("has_headset", "Casque"),
    ("shipping", "Livraison"),
    ("hand_delivery", "Remise en main propre"),
    ("posted_at", "Date publication"),
    ("age_text", "Âge annonce"),
    ("age_hours", "Âge (heures)"),
    ("best_deal_score", "Score affaire"),
    ("score_composite", "Score composite"),
    ("deal_tier", "Catégorie affaire"),
    ("spec_confidence", "Confiance specs"),
    ("analysis_detail", "Analyse"),
    ("enriched_by", "Enrichi par"),
]


def _cell_value(row: Dict[str, Any], key: str) -> Any:
    if key == "model":
        return row.get("model") or row.get("variant") or ""
    if key == "controllers_count":
        v = row.get("controllers_count")
        if v is None:
            v = row.get("controller_count")
        if v is None and row.get("has_controller") is False:
            return 0
        return v
    if key == "condition":
        return row.get("condition") or row.get("condition_hint") or ""
    if key == "price":
        return row.get("price") or row.get("total_price") or ""
    if key in ("games_list", "premium_games"):
        val = row.get(key)
        if isinstance(val, list):
            return ", ".join(str(x) for x in val)
        es = row.get("enriched_specs") or {}
        meta = es.get(key) or {}
        if isinstance(meta, dict) and meta.get("value"):
            v = meta["value"]
            return ", ".join(v) if isinstance(v, list) else v
        return val or ""
    if key == "hand_delivery":
        return row.get("hand_delivery") or row.get("hand_delivery_only") or ""
    if key == "shipping":
        return row.get("shipping") or row.get("delivery_mode") or ""
    return row.get(key, "")


def _auto_width(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = col[0].column_letter
        best = 10
        for cell in col:
            try:
                best = max(best, min(max_width, len(str(cell.value or "")) + 2))
            except Exception:
                pass
        ws.column_dimensions[letter].width = best


def _write_sheet(ws, rows: Sequence[Dict[str, Any]], columns: List[tuple[str, str]]) -> None:
    from openpyxl.styles import Font

    headers = [lbl for _, lbl in columns]
    keys = [k for k, _ in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_cell_value(row, k) for k in keys])
    ws.freeze_panes = "A2"
    _auto_width(ws)


def _market_stats(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    prices = []
    models: Dict[str, int] = {}
    for r in rows:
        try:
            p = float(r.get("total_price") or r.get("price") or 0)
        except (TypeError, ValueError):
            p = 0.0
        if p > 0:
            prices.append(p)
        mod = str(r.get("variant") or r.get("model") or "inconnu")
        models[mod] = models.get(mod, 0) + 1
    stats: Dict[str, Any] = {"count": len(prices), "models": models}
    if prices:
        sp = sorted(prices)
        n = len(sp)
        stats["p25"] = round(sp[max(0, n // 4)], 2)
        stats["p50"] = round(float(statistics.median(sp)), 2)
        stats["p75"] = round(sp[min(n - 1, (3 * n) // 4)], 2)
    return stats


def export_results_xlsx(
    path: Path,
    *,
    rows: List[Dict[str, Any]],
    top5: Optional[List[Dict[str, Any]]] = None,
    rejected: Optional[List[Dict[str, Any]]] = None,
    doubtful: Optional[List[Dict[str, Any]]] = None,
    category: str = "",
) -> Path:
    """Write results_enriched.xlsx with French sheets."""
    from openpyxl import Workbook

    path = Path(path)
    wb = Workbook()
    ws_offres = wb.active
    ws_offres.title = "Offres"
    _write_sheet(ws_offres, rows, OFFRES_COLUMNS)

    ws_top = wb.create_sheet("Top 5")
    top_rows = []
    for i, t in enumerate(top5 or [], 1):
        if not isinstance(t, dict):
            continue
        lid = str(t.get("listing_id") or t.get("id") or "")
        src = next((r for r in rows if str(r.get("id") or "") == lid), t)
        top_rows.append({**src, "top5_rank": i, "top5_score": t.get("score"), "top5_verdict": t.get("verdict")})
    top_cols = [
        ("top5_rank", "Rang"),
        ("title", "Titre"),
        ("price", "Prix"),
        ("best_deal_score", "Score affaire"),
        ("top5_score", "Score IA"),
        ("top5_verdict", "Verdict"),
        ("variant", "Modèle"),
        ("controllers_count", "Manettes"),
        ("games_count", "Jeux"),
        ("url", "URL"),
    ]
    _write_sheet(ws_top, top_rows or (rows[:5] if rows else []), top_cols)

    ws_marche = wb.create_sheet("Marché")
    stats = _market_stats(rows)
    ws_marche.append(["Indicateur", "Valeur"])
    ws_marche.append(["Catégorie", category])
    ws_marche.append(["Annonces", stats.get("count", 0)])
    for k in ("p25", "p50", "p75"):
        if k in stats:
            ws_marche.append([k.upper(), stats[k]])
    ws_marche.append([])
    ws_marche.append(["Modèle", "Nombre"])
    for mod, cnt in sorted((stats.get("models") or {}).items(), key=lambda x: -x[1]):
        ws_marche.append([mod, cnt])
    _auto_width(ws_marche)

    if rejected:
        ws_rej = wb.create_sheet("Rejetés")
        rej_cols = [
            ("id", "ID"),
            ("title", "Titre"),
            ("price", "Prix"),
            ("reject_reason", "Raison"),
            ("reject_stage", "Étape"),
            ("url", "URL"),
        ]
        _write_sheet(ws_rej, rejected, rej_cols)

    if doubtful:
        ws_dou = wb.create_sheet("Doutes")
        dou_cols = [
            ("id", "ID"),
            ("title", "Titre"),
            ("price", "Prix"),
            ("triage_reasons", "Raisons"),
            ("best_deal_score", "Score"),
            ("url", "URL"),
        ]
        _write_sheet(ws_dou, doubtful, dou_cols)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- report.py ---

import statistics
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional


@dataclass(slots=True)
class MarketReport:
    title: str = ""
    markdown: str = ""
    summary_text: str = ""
    stats: Dict[str, Any] = field(default_factory=dict)
    top_deals: List[Dict[str, Any]] = field(default_factory=list)


def _prices(rows: List[Dict[str, Any]]) -> List[float]:
    out: List[float] = []
    for r in rows:
        p = r.get("total_price") or r.get("price")
        try:
            v = float(p)
            if v > 0:
                out.append(v)
        except (TypeError, ValueError):
            continue
    return out


def _pct(x: float, lo: float, hi: float) -> float:
    if hi <= lo:
        return 0.0
    return max(0.0, min(100.0, (x - lo) / (hi - lo) * 100.0))


def build_market_report(
    rows: List[Dict[str, Any]],
    *,
    search_label: str = "",
    source_url: str = "",
    catalog_count: int = 0,
    excluded_accessories: int = 0,
    platform_analytics: Optional[Dict[str, Any]] = None,
) -> MarketReport:
    """Construit markdown + résumé GUI à partir des lignes scorées."""
    prices = _prices(rows)
    n = len(rows)
    rep = MarketReport(title=search_label or "Marché Vinted")

    if not prices:
        rep.summary_text = "Aucune annonce principale après filtrage (accessoires exclus)."
        rep.markdown = f"# Rapport — {rep.title}\n\n{rep.summary_text}\n"
        return rep

    lo, hi = min(prices), max(prices)
    med = statistics.median(prices)
    mean = statistics.mean(prices)
    try:
        p25 = statistics.quantiles(prices, n=4)[0]
        p75 = statistics.quantiles(prices, n=4)[2]
    except Exception:
        p25, p75 = lo, hi

    # Top deals for report table
    sorted_rows = sorted(rows, key=lambda r: float(r.get("final_score") or 0.0), reverse=True)
    top = sorted_rows[:12]
    rep.top_deals = top

    def _is_bad_flag(row: Dict[str, Any]) -> bool:
        if row.get("deal_blocked"):
            return True
        flags = (
            str(row.get("risk_flags") or "")
            + " "
            + str(row.get("ai_flags") or "")
            + " "
            + str(row.get("ai_spec_flags") or "")
        ).lower()
        bad = (
            "no_power",
            "broken",
            "for_parts",
            "shuts_off",
            "no_signal",
            "repair_only",
            "scam",
            "scam_risk",
            "ban",
            "banned",
            "icloud",
            "fake",
            "counterfeit",
        )
        return any(x in flags for x in bad)

    best_non_broken = next((r for r in sorted_rows if not _is_bad_flag(r)), sorted_rows[0])

    # Category breakdown
    by_cat: Dict[str, int] = {}
    for r in rows:
        c = str(r.get("item_type") or "unknown")
        by_cat[c] = by_cat.get(c, 0) + 1
    dominant_cat = max(by_cat.items(), key=lambda kv: kv[1])[0] if by_cat else "unknown"

    # Condition breakdown
    by_cond: Dict[str, int] = {}
    for r in rows:
        c = str(r.get("condition") or "—")[:40]
        by_cond[c] = by_cond.get(c, 0) + 1

    avg_disc = 0.0
    discs = []
    for r in rows:
        try:
            discs.append(float(r.get("market_discount_pct") or 0))
        except Exception:
            pass
    if discs:
        avg_disc = statistics.mean(discs)

    if platform_analytics is None:
        try:
            from .engine import compute_cross_platform_analytics

            platform_analytics = compute_cross_platform_analytics(rows, main_only=True)
        except Exception:
            platform_analytics = {}

    rep.stats = {
        "count": n,
        "catalog_count": catalog_count or n,
        "excluded_accessories": excluded_accessories,
        "min": lo,
        "max": hi,
        "median": med,
        "mean": mean,
        "p25": p25,
        "p75": p75,
        "avg_discount_pct": round(avg_disc, 1),
        "dominant_category": dominant_cat,
        "by_category": by_cat,
        "by_condition": by_cond,
        "platform_analytics": platform_analytics or {},
    }

    # Outliers (cheap + expensive)
    iqr = max(1e-6, float(p75 - p25))
    low_thr = p25 - 1.5 * iqr
    high_thr = p75 + 1.5 * iqr
    low_outliers = []
    high_outliers = []
    for r in rows:
        p = r.get("total_price") or r.get("price")
        try:
            pv = float(p)
        except Exception:
            continue
        if pv <= low_thr:
            low_outliers.append(r)
        if pv >= high_thr:
            high_outliers.append(r)
    low_outliers = sorted(low_outliers, key=lambda r: float(r.get("total_price") or r.get("price") or 0))[:5]
    high_outliers = sorted(high_outliers, key=lambda r: float(r.get("total_price") or r.get("price") or 0), reverse=True)[:5]
    rep.stats["outliers_low_thr"] = round(float(low_thr), 2)
    rep.stats["outliers_high_thr"] = round(float(high_thr), 2)
    rep.stats["outliers_low_count"] = len(low_outliers)
    rep.stats["outliers_high_count"] = len(high_outliers)
    rep.stats["outliers_low"] = [{"id": x.get("id"), "title": x.get("title"), "price": x.get("total_price") or x.get("price"), "url": x.get("url")} for x in low_outliers]
    rep.stats["outliers_high"] = [{"id": x.get("id"), "title": x.get("title"), "price": x.get("total_price") or x.get("price"), "url": x.get("url")} for x in high_outliers]

    # Summary for GUI (cool recap)
    lines = [
        f"🔎 {rep.title}",
        f"• {n} annonces principales analysées"
        + (f" (sur {catalog_count} brutes, {excluded_accessories} accessoires exclus)" if catalog_count else ""),
        f"• Prix médian : {med:.0f} €  |  fourchette {lo:.0f}–{hi:.0f} €",
        f"• Décote moyenne vs marché : {avg_disc:.0f}%",
        f"• Catégorie dominante : {dominant_cat}",
    ]
    pa = platform_analytics or {}
    if pa.get("platforms"):
        lines.append(f"• Sources : {pa.get('summary', '')}")
        for pk, ps in sorted(pa["platforms"].items(), key=lambda x: -x[1].get("count", 0)):
            lines.append(
                f"  - {ps.get('label', pk)} : {ps.get('count', 0)} offres ({ps.get('pct', 0):.0f}%) "
                f"· médiane {ps.get('median_eur', 0):.0f}€ · min {ps.get('min_eur', 0):.0f}€"
            )
    if top:
        best = best_non_broken
        lines.append(
            f"• Meilleure offre : score {float(best.get('final_score') or 0):.0f} — "
            f"{best.get('total_price') or best.get('price')} € — {(best.get('title') or '')[:50]}"
        )
        # Add a short \"Top 3\" preview
        lines.append("• Top 3 :")
        for r in top[:3]:
            lines.append(
                f"  - {float(r.get('final_score') or 0):.0f} | {r.get('total_price') or r.get('price')}€ | {(str(r.get('title') or '')[:60])}"
            )

    if low_outliers:
        lines.append("• Outliers bas (suspects / super deals) :")
        for r in low_outliers[:3]:
            lines.append(f"  - {r.get('price') or r.get('total_price')}€ | {(str(r.get('title') or '')[:55])}")
    rep.summary_text = "\n".join(lines)

    # Markdown report (export)
    md: List[str] = [
        f"# Rapport de marché — {rep.title}",
        "",
        f"- Source : `{source_url[:200]}{'…' if len(source_url) > 200 else ''}`",
        f"- Annonces principales : **{n}**"
        + (f" / {catalog_count} catalogue" if catalog_count else ""),
        f"- Accessoires exclus : **{excluded_accessories}**",
        "",
        "## Distribution des prix (EUR)",
        "",
        "| Min | P25 | Médiane | Moyenne | P75 | Max |",
        "|---:|---:|---:|---:|---:|---:|",
        f"| {lo:.0f} | {p25:.0f} | {med:.0f} | {mean:.0f} | {p75:.0f} | {hi:.0f} |",
        "",
        "## Top deals",
        "",
        "| Score | Décote | Prix | Titre | Cat | Lien |",
        "|---:|---:|---:|---|---|---|",
    ]
    for r in top:
        sc = r.get("final_score", "—")
        disc = r.get("market_discount_pct", "—")
        pr = r.get("total_price") or r.get("price")
        title = str(r.get("title") or "—").replace("|", " ")[:60]
        cat = r.get("item_type", "—")
        url = r.get("url", "")
        md.append(f"| {sc} | {disc}% | {pr} | {title} | {cat} | {url} |")

    md.extend(["", "## Par catégorie", "", "| Catégorie | Nb |", "|---|---:|"])
    for cat, cnt in sorted(by_cat.items(), key=lambda kv: -kv[1]):
        md.append(f"| {cat} | {cnt} |")

    md.extend(["", "## Par état", "", "| État | Nb |", "|---|---:|"])
    for cond, cnt in sorted(by_cond.items(), key=lambda kv: -kv[1])[:10]:
        md.append(f"| {cond} | {cnt} |")

    md.append("\n*Rapport généré par VF — version finale.*\n")
    rep.markdown = "\n".join(md)
    return rep


# --- eval_export.py ---

import json
import time
from pathlib import Path
from typing import Any, Dict, List, Optional


def _strip_null_bytes(val: Any) -> Any:
    if isinstance(val, str):
        return val.replace("\x00", "")
    if isinstance(val, dict):
        return {k: _strip_null_bytes(v) for k, v in val.items()}
    if isinstance(val, list):
        return [_strip_null_bytes(v) for v in val]
    if isinstance(val, tuple):
        return tuple(_strip_null_bytes(v) for v in val)
    return val


def _safe_json_write(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    clean = _strip_null_bytes(payload)
    path.write_text(json.dumps(clean, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _sample_deal(row: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "id": row.get("id"),
        "title": str(row.get("title") or "")[:120],
        "price": row.get("total_price") or row.get("price"),
        "platform": row.get("platform"),
        "final_score": row.get("final_score"),
        "score_composite": row.get("score_composite") or row.get("composite_score"),
        "deal_tier": row.get("deal_tier"),
        "rank_score": row.get("rank_score") or row.get("display_rank_score"),
        "url": row.get("url"),
        "analysis_detail": str(row.get("analysis_detail") or "")[:200],
    }


def write_run_eval_summary(
    *,
    run_folder,
    query: str,
    category: str,
    platform_selection: Optional[Dict[str, Any]] = None,
    scrape_stats: Optional[Dict[str, Dict[str, Any]]] = None,
    catalog_count: int = 0,
    main_count: int = 0,
    deals_count: int = 0,
    duration_s: float = 0.0,
    top_deals: Optional[List[Dict[str, Any]]] = None,
    dynamic_filters: Optional[List[Dict[str, Any]]] = None,
    top5_analysis: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    folder = Path(run_folder)
    eval_dir = folder / "eval"
    eval_dir.mkdir(parents=True, exist_ok=True)
    deals = list(top_deals or [])
    best = deals[0] if deals else {}
    cheapest_plat, cheapest_eur = "", 0.0
    if deals:
        cheapest = min(deals, key=lambda d: float(d.get("total_price") or d.get("price") or 1e9))
        cheapest_plat = str(cheapest.get("platform") or "")
        try:
            cheapest_eur = float(cheapest.get("total_price") or cheapest.get("price") or 0)
        except (TypeError, ValueError):
            pass
    summary: Dict[str, Any] = {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "query": query,
        "category": category,
        "duration_s": round(float(duration_s or 0), 1),
        "catalog_count": catalog_count,
        "main_items": main_count,
        "deals_exported": deals_count,
        "platform_selection": platform_selection or {},
        "scrape_stats": scrape_stats or {},
        "dynamic_filter_ids": [str(f.get("id") or "") for f in (dynamic_filters or []) if f.get("id")],
        "highlights": {
            "best_deal_title": str(best.get("title") or "")[:100],
            "best_deal_price_eur": best.get("total_price") or best.get("price"),
            "best_deal_platform": best.get("platform"),
            "cheapest_among_top_platform": cheapest_plat,
            "cheapest_among_top_eur": cheapest_eur,
        },
        "top_samples": [_sample_deal(d) for d in deals[:8]],
        "top5_ia": top5_analysis.get("top5") if isinstance(top5_analysis, dict) else None,
        "ai_review_prompt": "Review this VF Finder run: triage, filters, platform selection, scoring.",
    }
    summary_path = eval_dir / "eval_summary.json"
    _safe_json_write(summary_path, summary)
    prompt_path = eval_dir / "eval_prompt.md"
    prompt_path.write_text(f"# Eval run\n\n- Requete: `{query}`\n- Categorie: `{category}`\n", encoding="utf-8")
    return {"eval_summary": str(summary_path.resolve()), "eval_prompt": str(prompt_path.resolve())}


# --- market_report.py ---
from typing import Any, Dict, List, Optional


def _fmt_eur(x: Any) -> str:
    try:
        v = float(x)
        if v <= 0:
            return "-"
        if abs(v - round(v)) < 1e-9:
            return f"{v:.0f} EUR"
        return f"{v:.2f} EUR"
    except Exception:
        return "-"


def _safe_int(x: Any) -> Optional[int]:
    try:
        return int(float(x))
    except Exception:
        return None


def build_analysis_markdown(
    *,
    category: str = "",
    query: str = "",
    brain: Optional[Dict[str, Any]] = None,
    rows: Optional[List[Dict[str, Any]]] = None,
    platform_analytics: Optional[Dict[str, Any]] = None,
    triage_stats: Optional[Dict[str, Any]] = None,
) -> str:
    """
    GUI-facing analysis markdown.

    Intentionally lightweight and robust: should never raise, and should still
    provide useful context even when some pipeline steps failed.
    """
    rows = list(rows or [])
    cat = (category or "").strip() or "generic"
    q = (query or "").strip()

    n_rows = len(rows)
    n_doubt = sum(1 for r in rows if str(r.get("triage_decision") or "") == "doubtful")
    n_acc = sum(1 for r in rows if r.get("is_accessory"))

    prices: List[float] = []
    for r in rows:
        p = r.get("total_price") or r.get("price")
        try:
            v = float(p)
            if v > 0:
                prices.append(v)
        except Exception:
            continue
    prices.sort()
    p25 = prices[int(len(prices) * 0.25)] if prices else None
    p50 = prices[int(len(prices) * 0.50)] if prices else None
    p75 = prices[int(len(prices) * 0.75)] if prices else None

    lines: List[str] = []
    lines += ["## Analyse", ""]
    if q:
        tail = "..." if len(q) > 220 else ""
        lines.append(f"- Requete : `{q[:220]}{tail}`")
    lines.append(f"- Marche : **{cat}**")
    if triage_stats:
        sc = _safe_int(triage_stats.get("scraped"))
        at = _safe_int(triage_stats.get("after_triage"))
        rj = _safe_int(triage_stats.get("rejected"))
        db = _safe_int(triage_stats.get("doubts"))
        bits: List[str] = []
        if sc is not None:
            bits.append(f"scrapees {sc}")
        if at is not None:
            bits.append(f"retenues {at}")
        if rj is not None:
            bits.append(f"rejetees {rj}")
        if db is not None:
            bits.append(f"doutes {db}")
        if bits:
            lines.append(f"- Pipeline : {' | '.join(bits)}")
    lines.append(f"- Lignes resultats (toutes) : **{n_rows}**")
    if n_acc:
        lines.append(f"- Accessoires reperes : **{n_acc}**")
    if n_doubt:
        lines.append(f"- Doutes (triage) : **{n_doubt}**")
    if prices:
        lines.append(
            f"- Prix (approx.) : P25 {_fmt_eur(p25)} | mediane {_fmt_eur(p50)} | P75 {_fmt_eur(p75)}"
        )
    lines.append("")

    if platform_analytics and platform_analytics.get("summary"):
        lines += ["## Plateformes", "", str(platform_analytics.get("summary") or "-"), ""]

    if isinstance(brain, dict) and brain:
        lines += ["## Signaux (echantillon)", ""]
        ref = brain.get("market_reference") if isinstance(brain.get("market_reference"), dict) else {}
        if ref:
            if ref.get("median"):
                lines.append(f"- Reference marche (sample) : mediane {_fmt_eur(ref.get('median'))}")
            if ref.get("count"):
                lines.append(f"- Taille echantillon : {ref.get('count')}")
        note = str(brain.get("note") or brain.get("summary") or "").strip()
        if note:
            lines.append(f"- Note : {note[:400]}")
        lines.append("")

    lines += [
        "## Lecture rapide",
        "",
        "- Onglet **Selection IA** : Top 5 IA (si dispo) + Top affaires (regles).",
        "- Onglet **Doutes** : annonces ambigues (titre/etat/pack) a trier.",
        "- Onglet **Rejetes** : hors-sujet / accessoires / filtrage strict (archive).",
    ]
    return "\n".join(lines).strip() + "\n"

