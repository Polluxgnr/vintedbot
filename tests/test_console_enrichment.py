"""Unit tests — PS4 console spec enrichment."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from vf.scoring import (
    ConsolePS4Enricher,
    detect_premium_games,
    extract_console_ps4_specs,
    premium_games_boost,
)
from vf.scoring import enrich_listings_specs


class ConsoleEnrichmentTests(unittest.TestCase):
    def test_extract_slim_2_manettes_jeux(self):
        listing = {
            "title": "PS4 Slim 1To + 2 manettes + FIFA 23 + GTA V",
            "description": "Très bon état, boîte, HDMI et alimentation",
        }
        specs = extract_console_ps4_specs(listing)
        self.assertEqual(specs["variant"]["value"], "slim")
        self.assertEqual(specs["storage_gb"]["value"], 1000)
        self.assertEqual(specs["controllers_count"]["value"], 2)
        self.assertGreaterEqual(specs["games_count"]["value"], 2)
        premium = specs["premium_games"]["value"]
        self.assertIn("fifa", premium)
        self.assertIn("gta", premium)

    def test_sans_manette_penalty_fields(self):
        listing = {"title": "PS4 slim noir sans manette", "description": "Console seule 500 Go"}
        specs = extract_console_ps4_specs(listing)
        self.assertEqual(specs["controllers_count"]["value"], 0)
        self.assertFalse(specs["has_controller"]["value"])

    def test_enricher_flattens_row(self):
        listing = {
            "id": "42",
            "title": "PS4 Pro 1To 2 manettes God of War",
            "description": "",
        }
        ConsolePS4Enricher().apply_to_listing(listing, use_ai=False)
        self.assertEqual(listing["variant"], "pro")
        self.assertEqual(listing["controllers_count"], 2)
        self.assertIn("god_of_war", listing["premium_games"])
        self.assertIn("enriched_specs", listing)
        self.assertGreater(listing["spec_confidence"], 0.5)

    def test_premium_boost_capped(self):
        keys = list(detect_premium_games("FIFA GTA God of War Last of Us Spider-Man Horizon"))
        boost = premium_games_boost(keys)
        self.assertGreater(boost, 20.0)
        self.assertLessEqual(boost, 35.0)

    def test_batch_enrich_listing_specs(self):
        rows = [
            {"id": "1", "title": "PS4 + manette", "price": 90},
            {"id": "2", "title": "PS4 slim 500go 2 manettes FIFA", "price": 120},
        ]
        n = enrich_listings_specs(rows, "console_ps4", use_ai=False, limit=10)
        self.assertEqual(n, 2)
        self.assertTrue(rows[1].get("controllers_count") >= 2)

    def test_excel_export_columns(self):
        try:
            from openpyxl import load_workbook
            from vf.export import OFFRES_COLUMNS, export_results_xlsx
        except ImportError:
            self.skipTest("openpyxl not installed")

        row = {
            "id": "99",
            "title": "PS4 Slim pack complet",
            "price": 110,
            "platform": "vinted",
            "variant": "slim",
            "controllers_count": 2,
            "games_count": 3,
            "games_list": ["FIFA 23"],
            "premium_games": ["fifa"],
            "best_deal_score": 78.5,
            "enriched_by": "rules",
        }
        with tempfile.TemporaryDirectory() as td:
            path = export_results_xlsx(
                Path(td) / "results_enriched.xlsx",
                rows=[row],
                category="console_ps4",
            )
            wb = load_workbook(path)
            ws = wb["Offres"]
            headers = [c.value for c in ws[1]]
            self.assertEqual(headers[0], "ID")
            self.assertIn("Manettes", headers)
            self.assertEqual(ws[2][0].value, "99")


if __name__ == "__main__":
    unittest.main()
